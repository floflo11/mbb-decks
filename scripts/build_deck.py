#!/usr/bin/env python3
"""
build_deck.py — Generate an MBB-style consulting deck (.pptx) from a JSON spec.

Usage:
    python scripts/build_deck.py <input.json> <output.pptx>

The JSON spec defines the six-section MBB structure:
    cover -> executive_summary -> agenda -> sections -> recommendations -> appendix

See examples/market-entry/input.json for the canonical schema.
"""

from __future__ import annotations

import json
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

# Friendly error if python-pptx is missing, since this is by far the
# most common first-run failure for new users.
if sys.version_info < (3, 9):
    print(
        f"ERROR: mbb-decks needs Python 3.9 or newer. You have {sys.version.split()[0]}.\n"
        "Install a newer Python from https://www.python.org/downloads/ or via your package manager.",
        file=sys.stderr,
    )
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError as e:
    print(
        f"\nERROR: mbb-decks needs python-pptx, which is not installed.\n"
        f"  Missing module: {e.name}\n\n"
        "Install it with:\n"
        "    pip install python-pptx\n\n"
        "Or, if you use a different Python install:\n"
        "    python3 -m pip install python-pptx\n",
        file=sys.stderr,
    )
    sys.exit(1)


# Visual system tokens (see reference/visual-system.md)
NAVY = RGBColor(0x05, 0x1C, 0x2C)
NEAR_BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BLACK = RGBColor(0x00, 0x00, 0x00)
LIGHT_GREY = RGBColor(0xE5, 0xE7, 0xEB)
MID_GREY = RGBColor(0x94, 0x9B, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BLUE = RGBColor(0x22, 0x51, 0xFF)

HEADLINE_FONT = "Georgia"
BODY_FONT = "Calibri"

CHART_PALETTE = [NAVY, ACCENT_BLUE, MID_GREY, LIGHT_GREY]

# Slide dimensions (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Layout constants — every action-titled slide uses the same template
M_LEFT = Inches(0.5)
M_RIGHT = Inches(0.5)
TITLE_Y = Inches(0.4)
TITLE_H = Inches(0.7)        # one line of headline at up to 22pt Georgia
RULE_Y = Inches(1.15)        # SAME rule position on every action-titled slide
BODY_Y = Inches(1.35)
BODY_H = Inches(5.25)
FOOTNOTE_Y = Inches(6.7)
SOURCE_Y = Inches(7.0)
PAGE_NUM_Y = Inches(7.0)

HEADLINE_MAX_CHARS = 110     # warn if action title would wrap past one line

# Logo cache for inline bullet markers (auto-fetched from Hunter.io)
LOGO_CACHE_DIR = Path(__file__).resolve().parent.parent / "assets" / "logos"
HUNTER_LOGO_URL = "https://logos.hunter.io/{domain}"


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _add_text(slide, x, y, w, h, text, *, font=BODY_FONT, size=10, color=NEAR_BLACK,
              bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return box


def _add_headline(slide, text, size=18):
    """Action title or headline. Always one line, always at TITLE_Y, rule below at RULE_Y."""
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    return _add_text(
        slide, M_LEFT, TITLE_Y, body_w, TITLE_H, text,
        font=HEADLINE_FONT, size=size, color=NAVY,
        anchor=MSO_ANCHOR.TOP, wrap=False,
    )


def _add_rule(slide, y=RULE_Y, color=BLACK, weight=0.75):
    line = slide.shapes.add_connector(1, M_LEFT, y, SLIDE_W - M_RIGHT, y)
    line.line.color.rgb = color
    line.line.width = Pt(weight)
    return line


def _add_bullets(slide, x, y, w, h, bullets, *, size=14, color=NEAR_BLACK, line_spacing=1.3):
    """
    Render a bullet list. Each item is either:
      - a string (plain "•" marker), or
      - {"icon": glyph, "text": str} (small navy circle with glyph as marker), or
      - {"logo": domain, "text": str} (company logo as marker; auto-fetched).
    If both "icon" and "logo" are present, "logo" wins.
    """
    if not bullets:
        return None
    has_markers = any(
        isinstance(b, dict) and (b.get("icon") or b.get("logo"))
        for b in bullets
    )
    if has_markers:
        return _add_iconed_bullets(slide, x, y, w, h, bullets, size=size, color=color)
    return _add_plain_bullets(slide, x, y, w, h, bullets, size=size, color=color,
                              line_spacing=line_spacing)


def _add_plain_bullets(slide, x, y, w, h, bullets, *, size=14, color=NEAR_BLACK, line_spacing=1.3):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        text = item if isinstance(item, str) else item.get("text", "")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(10)

        marker = p.add_run()
        marker.text = "•   "
        marker.font.name = BODY_FONT
        marker.font.size = Pt(size)
        marker.font.color.rgb = NAVY

        body = p.add_run()
        body.text = text
        body.font.name = BODY_FONT
        body.font.size = Pt(size)
        body.font.color.rgb = color
    return box


def _add_small_icon(slide, x, y, glyph, size):
    """Small navy stroke-only circle with a Unicode glyph centered inside."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = NAVY
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = glyph
    run.font.name = HEADLINE_FONT
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    return shape


def _add_inline_logo(slide, x, y, max_w, max_h, logo_path):
    """Place a company logo as an inline marker, preserving aspect ratio.

    Target height is capped to keep all logos visually consistent across bullets;
    width adapts to aspect ratio but is clamped to max_w so the logo never spills
    into the text column.
    """
    iw, ih = _png_dimensions(logo_path)
    aspect = (iw / ih) if (iw and ih) else 1.0
    target_h = min(Inches(0.45), int(max_h * 0.65))
    target_w = int(target_h * aspect)
    if target_w > max_w:
        target_w = int(max_w)
        target_h = int(target_w / aspect)
    cx = int(x + (max_w - target_w) / 2)
    cy = int(y + (max_h - target_h) / 2)
    try:
        slide.shapes.add_picture(str(logo_path), cx, cy,
                                 width=target_w, height=target_h)
    except Exception as e:
        print(f"WARNING: failed to embed inline logo {logo_path}: {e}", file=sys.stderr)


def _add_iconed_bullets(slide, x, y, w, h, bullets, *, size=14, color=NEAR_BLACK):
    """Per-bullet marker layout. Marker is logo if 'logo' set, else navy-circle icon if 'icon' set, else plain bullet."""
    n = len(bullets)
    icon_size = Inches(0.7)              # wider so wide logos render legibly
    icon_pad = Inches(0.2)
    row_h = int(min(h / n, Inches(1.2)))
    text_x = x + icon_size + icon_pad
    text_w = w - (icon_size + icon_pad)

    for i, item in enumerate(bullets):
        row_y = y + i * row_h
        if isinstance(item, dict):
            glyph = item.get("icon", "")
            domain = item.get("logo")
            text = item.get("text", "")
        else:
            glyph = ""
            domain = None
            text = item

        logo_path = _ensure_logo(domain) if domain else None
        if logo_path:
            _add_inline_logo(slide, x, row_y, icon_size, row_h, logo_path)
        elif glyph:
            # Plain navy glyph (no circle) so abstract icons sit at the same visual
            # weight as the company logos in the same column.
            _add_text(slide, x, row_y, icon_size, row_h, glyph,
                      font=HEADLINE_FONT, size=22, color=NAVY,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            _add_text(slide, x, row_y, icon_size, row_h, "•",
                      font=BODY_FONT, size=size, color=NAVY,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _add_text(slide, text_x, row_y, text_w, row_h, text,
                  font=BODY_FONT, size=size, color=color, anchor=MSO_ANCHOR.MIDDLE, wrap=True)


_LINK_PATTERN = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
_url_404_cache = {}


def _http_status(url, method, timeout=4):
    try:
        req = urllib.request.Request(
            url, method=method,
            headers={"User-Agent": "Mozilla/5.0 (compatible; mbb-decks/1.0)"},
        )
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status
    except urllib.error.HTTPError as e:
        return e.code
    except Exception:
        return None  # network error / DNS / timeout


def _is_url_404(url, timeout=4):
    """Return True only if the URL definitively returns 404 or 410.

    HEAD-rejecting sites (405, 403) are retried with GET so we don't drop
    valid links. Network errors and timeouts are treated as 'unverifiable'
    and the link is preserved. Result is cached.
    """
    if url in _url_404_cache:
        return _url_404_cache[url]
    status = _http_status(url, "HEAD", timeout)
    # If HEAD is rejected (405) or auth-gated (401, 403), retry with GET
    if status in (401, 403, 405):
        status = _http_status(url, "GET", timeout)
    broken = status in (404, 410)
    _url_404_cache[url] = broken
    if broken:
        print(f"NOTE: source URL hidden (HTTP {status}): {url}", file=sys.stderr)
    return broken


def _parse_source_segments(text):
    """Yield (segment_text, url_or_None) tuples by splitting on [text](url)."""
    pos = 0
    for m in _LINK_PATTERN.finditer(text):
        if m.start() > pos:
            yield (text[pos:m.start()], None)
        yield (m.group(1), m.group(2))
        pos = m.end()
    if pos < len(text):
        yield (text[pos:], None)


def _add_source_line(slide, x, y, w, source_text):
    """Render the source line with markdown-style [text](url) becoming hyperlinks."""
    box = slide.shapes.add_textbox(x, y, w, Inches(0.3))
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT

    # "Source: " prefix
    prefix = p.add_run()
    prefix.text = "Source: "
    prefix.font.name = BODY_FONT
    prefix.font.size = Pt(8)
    prefix.font.color.rgb = MID_GREY

    for segment, url in _parse_source_segments(source_text):
        if not segment:
            continue
        run = p.add_run()
        run.text = segment
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        if url and not _is_url_404(url):
            run.hyperlink.address = url
            # Subtle visual cue that this segment is clickable
            run.font.color.rgb = NEAR_BLACK
            run.font.underline = True
        else:
            run.font.color.rgb = MID_GREY


def _add_footer(slide, page_num, total, *, footnote=None, source=None, show_page=True):
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    if footnote:
        # Prefix with `* ` to mark it as a footnote. The body text or chart
        # label that the note refers to should carry a matching `*` marker.
        text = footnote if footnote.lstrip().startswith("*") else f"* {footnote}"
        _add_text(slide, M_LEFT, FOOTNOTE_Y, body_w, Inches(0.25),
                  text, size=8, color=MID_GREY, italic=True)
    if source:
        _add_source_line(slide, M_LEFT, SOURCE_Y, int(body_w * 0.75), source)
    if show_page:
        _add_text(slide, SLIDE_W - Inches(2.0) - M_RIGHT, PAGE_NUM_Y,
                  Inches(2.0), Inches(0.3),
                  f"Page {page_num} / {total}",
                  size=8, color=MID_GREY, align=PP_ALIGN.RIGHT)


def _style_chart(chart, *, palette=CHART_PALETTE):
    """MBB chart style: data labels on every bar, no gridlines, minimal axes."""
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name = BODY_FONT
    chart.legend.font.size = Pt(10)

    for i, series in enumerate(chart.series):
        fill = series.format.fill
        fill.solid()
        fill.fore_color.rgb = palette[i % len(palette)]
        series.format.line.fill.background()

    # Data labels carry the precision. Place above each bar.
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.name = BODY_FONT
    dl.font.size = Pt(9)
    dl.font.color.rgb = NEAR_BLACK
    dl.number_format = "#,##0.0"   # 1 decimal place across the entire deck

    # Y-axis: drop gridlines and ticks; suppress the numeric labels
    # since the data labels above each bar already show the values.
    va = chart.value_axis
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.number_format = ";;;"          # suppress all numeric labels
    va.tick_labels.font.size = Pt(1)

    # X-axis: keep category labels, drop ticks
    ca = chart.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.name = BODY_FONT
    ca.tick_labels.font.size = Pt(10)
    ca.tick_labels.font.color.rgb = NEAR_BLACK


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def build_cover(prs, meta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    # Thin accent rule near the top
    rule = slide.shapes.add_connector(1, M_LEFT, Inches(2.3),
                                      M_LEFT + Inches(0.8), Inches(2.3))
    rule.line.color.rgb = ACCENT_BLUE
    rule.line.width = Pt(2.5)

    _add_text(slide, M_LEFT, Inches(2.6), body_w, Inches(1.6),
              meta.get("title", "Untitled"),
              font=HEADLINE_FONT, size=40, color=NAVY)

    if meta.get("subtitle"):
        _add_text(slide, M_LEFT, Inches(4.3), body_w, Inches(0.6),
                  meta["subtitle"],
                  font=HEADLINE_FONT, size=18, color=NEAR_BLACK, italic=True)

    byline = " | ".join(s for s in [meta.get("author"), meta.get("organization")] if s)
    if byline:
        _add_text(slide, M_LEFT, Inches(6.85), body_w, Inches(0.3),
                  byline, size=10, color=MID_GREY)
    return slide


def build_executive_summary(prs, summary, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    _add_headline(slide, summary["headline"], size=18)
    _add_rule(slide)
    _add_bullets(slide, M_LEFT, BODY_Y, body_w, BODY_H,
                 summary.get("bullets", []), size=14)
    _add_footer(slide, page, total)
    return slide


def build_agenda(prs, agenda, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    _add_headline(slide, "Agenda", size=18)
    _add_rule(slide)

    box = slide.shapes.add_textbox(M_LEFT, BODY_Y, body_w, BODY_H)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)

    for i, item in enumerate(agenda):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(18)

        num = p.add_run()
        num.text = f"0{i + 1}    "
        num.font.name = HEADLINE_FONT
        num.font.size = Pt(28)
        num.font.color.rgb = ACCENT_BLUE

        title = p.add_run()
        title.text = item["title"]
        title.font.name = HEADLINE_FONT
        title.font.size = Pt(22)
        title.font.color.rgb = NAVY

        if item.get("subtitle"):
            p2 = tf.add_paragraph()
            p2.line_spacing = 1.0
            p2.space_after = Pt(14)
            sub = p2.add_run()
            sub.text = "         " + item["subtitle"]
            sub.font.name = BODY_FONT
            sub.font.size = Pt(12)
            sub.font.color.rgb = MID_GREY
            sub.font.italic = True
    _add_footer(slide, page, total)
    return slide


def build_section_divider(prs, idx, title, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    rule = slide.shapes.add_connector(1, M_LEFT, Inches(3.1),
                                      M_LEFT + Inches(0.8), Inches(3.1))
    rule.line.color.rgb = ACCENT_BLUE
    rule.line.width = Pt(2.5)
    _add_text(slide, M_LEFT, Inches(3.3), body_w, Inches(0.7),
              f"0{idx}", font=HEADLINE_FONT, size=44, color=ACCENT_BLUE)
    _add_text(slide, M_LEFT, Inches(4.1), body_w, Inches(1.0),
              title, font=HEADLINE_FONT, size=32, color=NAVY)
    _add_footer(slide, page, total)
    return slide


def build_content_slide(prs, spec, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    _add_headline(slide, spec["action_title"], size=18)
    _add_rule(slide)

    content_type = spec.get("content_type", "bullets")
    if content_type == "bullets":
        _add_bullets(slide, M_LEFT, BODY_Y, body_w, BODY_H,
                     spec.get("bullets", []), size=14)
    elif content_type == "bar_chart":
        if spec.get("commentary"):
            _add_bar_chart_with_commentary(
                slide, spec["chart"], spec["commentary"],
                spec.get("commentary_position", "right"),
            )
        else:
            _add_bar_chart(slide, spec["chart"])
    elif content_type == "two_column":
        _add_two_column(
            slide,
            spec.get("left", []), spec.get("right", []),
            spec.get("left_header"), spec.get("right_header"),
            spec.get("left_icon"), spec.get("right_icon"),
        )

    _add_footer(slide, page, total,
                footnote=spec.get("footnote"),
                source=spec.get("source"))
    return slide


def _png_dimensions(path):
    """Read width, height from PNG header bytes; return (None, None) on failure."""
    try:
        with open(path, "rb") as f:
            header = f.read(24)
        if header[:8] != b"\x89PNG\r\n\x1a\n":
            return None, None
        return int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")
    except Exception:
        return None, None


def _ensure_logo(domain):
    """Return Path to cached logo, downloading via Hunter.io if missing. None on failure."""
    if not domain:
        return None
    LOGO_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    safe = domain.strip().lower().replace("/", "_")
    out = LOGO_CACHE_DIR / f"{safe}.png"
    if out.exists():
        return out
    url = HUNTER_LOGO_URL.format(domain=domain)
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "mbb-decks/1.0"})
        with urllib.request.urlopen(req, timeout=10) as resp:
            if resp.status != 200:
                return None
            data = resp.read()
        if len(data) < 200 or not data.startswith(b"\x89PNG"):
            return None
        out.write_bytes(data)
        return out
    except Exception as e:
        print(f"WARNING: logo fetch failed for {domain}: {e}", file=sys.stderr)
        return None


def _render_bar_chart(slide, chart_spec, x, y, w, h):
    """Render a clustered column chart inside a fixed bounding box."""
    cd = CategoryChartData()
    cd.categories = chart_spec["categories"]
    for series in chart_spec["series"]:
        cd.add_series(series["name"], series["values"])
    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd,
    )
    chart = chart_frame.chart
    _style_chart(chart)
    all_values = [v for s in chart_spec["series"] for v in s["values"]]
    max_v, min_v = max(all_values), min(all_values)
    chart.value_axis.maximum_scale = max_v * 1.15 if max_v > 0 else max_v * 0.85
    if min_v < 0:
        chart.value_axis.minimum_scale = min_v * 1.20
    return chart


def _add_unit_label(slide, x, y, w, text):
    _add_text(slide, x, y, w, Inches(0.25),
              text, size=9, color=MID_GREY, italic=True)


def _add_commentary_header(slide, x, y, w):
    _add_text(slide, x, y, w, Inches(0.3),
              "KEY TAKEAWAYS", size=9, color=MID_GREY, bold=True)


def _add_bar_chart(slide, chart_spec):
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    if chart_spec.get("unit"):
        _add_unit_label(slide, M_LEFT, BODY_Y - Inches(0.05), Inches(8), chart_spec["unit"])
    _render_bar_chart(slide, chart_spec,
                      M_LEFT, BODY_Y + Inches(0.2), body_w, Inches(4.3))


def _add_bar_chart_with_commentary(slide, chart_spec, commentary, position="right"):
    """Chart + commentary panel on the same slide. position: 'right' or 'below'."""
    body_w = SLIDE_W - M_LEFT - M_RIGHT

    if position == "below":
        # Reserve enough room for commentary so it does not overlap the footnote/source.
        chart_x = M_LEFT
        chart_w = body_w
        chart_y = BODY_Y + Inches(0.2)
        chart_h = Inches(2.8)            # was 3.4; commentary needs more room
        comm_x = M_LEFT
        comm_y = chart_y + chart_h + Inches(0.2)
        comm_w = body_w
        comm_size = 10                   # was 11; tighter for 2-column below layout
    else:  # "right"
        # Wrap float math in int() to keep EMU integer; PowerPoint rejects fractional EMUs.
        chart_w = int(body_w * 0.62)
        gap = Inches(0.35)
        chart_x = M_LEFT
        chart_y = BODY_Y + Inches(0.2)
        chart_h = Inches(4.3)
        comm_x = M_LEFT + chart_w + gap
        comm_y = BODY_Y
        comm_w = int(body_w - chart_w - gap)
        comm_size = 12

    if chart_spec.get("unit"):
        _add_unit_label(slide, chart_x, BODY_Y - Inches(0.05), chart_w, chart_spec["unit"])
    _render_bar_chart(slide, chart_spec, chart_x, chart_y, chart_w, chart_h)

    _add_commentary_header(slide, comm_x, comm_y, comm_w)
    bullets_y = comm_y + Inches(0.35)
    # Cap commentary height so it never overlaps the footnote line at FOOTNOTE_Y.
    bullets_h_max = FOOTNOTE_Y - bullets_y - Inches(0.1)
    if position == "below" and len(commentary) >= 4:
        mid = (len(commentary) + 1) // 2
        col_w = int((comm_w - Inches(0.4)) / 2)
        _add_bullets(slide, comm_x, bullets_y, col_w, bullets_h_max,
                     commentary[:mid], size=comm_size)
        _add_bullets(slide, comm_x + col_w + Inches(0.4), bullets_y,
                     col_w, bullets_h_max, commentary[mid:], size=comm_size)
    else:
        h = bullets_h_max if position == "below" else Inches(4.3)
        _add_bullets(slide, comm_x, bullets_y, comm_w, h,
                     commentary, size=comm_size)


def _add_column_icon(slide, x, y, glyph, size=Inches(0.55)):
    """Minimalist navy stroke-only circle with a Unicode glyph inside."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = NAVY
    shape.line.width = Pt(1.0)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = glyph
    run.font.name = HEADLINE_FONT
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY
    return shape


def _add_two_column(slide, left, right, left_header=None, right_header=None,
                    left_icon=None, right_icon=None):
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    col_w = int((body_w - Inches(0.4)) / 2)
    icon_size = Inches(0.55)
    icon_pad = Inches(0.18)
    header_h = Inches(0.55)
    bullets_top = BODY_Y + icon_size + Inches(0.25)

    def _render_column(x, header_text, icon_glyph, bullets):
        text_x = x
        text_w = col_w
        if icon_glyph:
            _add_column_icon(slide, x, BODY_Y, icon_glyph, icon_size)
            text_x = x + icon_size + icon_pad
            text_w = col_w - (icon_size + icon_pad)
        if header_text:
            _add_text(slide, text_x, BODY_Y + Inches(0.05), text_w, header_h,
                      header_text, font=HEADLINE_FONT, size=15, color=NAVY,
                      anchor=MSO_ANCHOR.TOP)
        _add_bullets(slide, x, bullets_top, col_w, BODY_H - (bullets_top - BODY_Y),
                     bullets, size=13)

    _render_column(M_LEFT, left_header, left_icon, left)
    _render_column(M_LEFT + col_w + Inches(0.4), right_header, right_icon, right)


def build_recommendations(prs, rec, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    _add_headline(slide, rec["headline"], size=18)
    _add_rule(slide)

    actions = rec.get("actions", [])
    if actions:
        rows = len(actions) + 1
        h = Inches(min(0.55 * rows + 0.2, 4.0))
        table_shape = slide.shapes.add_table(rows, 3, M_LEFT, BODY_Y, body_w, h)
        table = table_shape.table
        table.columns[0].width = Inches(2.0)
        table.columns[1].width = Inches(7.0)
        table.columns[2].width = body_w - Inches(9.0)

        for c, header in enumerate(["Owner", "Action", "Outcome"]):
            cell = table.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = header
            run.font.name = BODY_FONT
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = WHITE

        for r, action in enumerate(actions, start=1):
            for c, key in enumerate(["owner", "action", "outcome"]):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = action.get(key, "")
                run.font.name = BODY_FONT
                run.font.size = Pt(11)
                run.font.color.rgb = NEAR_BLACK
    _add_footer(slide, page, total)
    return slide


def build_appendix_divider(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body_w = SLIDE_W - M_LEFT - M_RIGHT
    _add_text(slide, M_LEFT, Inches(3.3), body_w, Inches(1.2),
              "Appendix", font=HEADLINE_FONT, size=44, color=NAVY)
    _add_text(slide, M_LEFT, Inches(4.3), body_w, Inches(0.4),
              "Backup detail and supporting analysis",
              size=14, color=MID_GREY, italic=True)
    _add_footer(slide, page, total)
    return slide


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def count_main_slides(spec):
    n = 1  # cover
    if "executive_summary" in spec:
        n += 1
    if "agenda" in spec:
        n += 1
    for section in spec.get("sections", []):
        n += 1 + len(section.get("slides", []))
    if "recommendations" in spec:
        n += 1
    return n


def count_total(spec):
    n = count_main_slides(spec)
    if spec.get("appendix"):
        n += 1 + len(spec["appendix"])
    return n


def _check_headline_lengths(spec):
    long_ones = []
    if "executive_summary" in spec:
        h = spec["executive_summary"].get("headline", "")
        if len(h) > HEADLINE_MAX_CHARS:
            long_ones.append(("executive_summary", h))
    for i, section in enumerate(spec.get("sections", []), 1):
        for j, slide in enumerate(section.get("slides", []), 1):
            t = slide.get("action_title", "")
            if len(t) > HEADLINE_MAX_CHARS:
                long_ones.append((f"section {i}, slide {j}", t))
    if "recommendations" in spec:
        h = spec["recommendations"].get("headline", "")
        if len(h) > HEADLINE_MAX_CHARS:
            long_ones.append(("recommendations", h))
    for j, slide in enumerate(spec.get("appendix", []), 1):
        t = slide.get("action_title", "")
        if len(t) > HEADLINE_MAX_CHARS:
            long_ones.append((f"appendix slide {j}", t))
    for where, text in long_ones:
        print(
            f"WARNING: headline at {where} is {len(text)} chars (max {HEADLINE_MAX_CHARS}); "
            f"will overflow one line: {text!r}",
            file=sys.stderr,
        )


def main(input_path, output_path):
    spec = json.loads(Path(input_path).read_text())

    _check_headline_lengths(spec)

    main_count = count_main_slides(spec)
    if main_count > 15:
        print(
            f"WARNING: main deck has {main_count} slides; MBB convention caps at 15. "
            "Move detail to appendix.",
            file=sys.stderr,
        )

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total = count_total(spec)
    page = 1

    build_cover(prs, spec.get("meta", {}))
    page += 1

    if "executive_summary" in spec:
        build_executive_summary(prs, spec["executive_summary"], page, total)
        page += 1
    if "agenda" in spec:
        build_agenda(prs, spec["agenda"], page, total)
        page += 1

    for i, section in enumerate(spec.get("sections", []), start=1):
        build_section_divider(prs, i, section["title"], page, total)
        page += 1
        for slide_spec in section.get("slides", []):
            build_content_slide(prs, slide_spec, page, total)
            page += 1

    if "recommendations" in spec:
        build_recommendations(prs, spec["recommendations"], page, total)
        page += 1

    if spec.get("appendix"):
        build_appendix_divider(prs, page, total)
        page += 1
        for slide_spec in spec["appendix"]:
            build_content_slide(prs, slide_spec, page, total)
            page += 1

    prs.save(output_path)
    print(f"Wrote {output_path} ({total} slides)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_deck.py <input.json> <output.pptx>", file=sys.stderr)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
